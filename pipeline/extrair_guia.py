"""Extrai do "Guia de Fundos" da XP (planilha .xlsx) a lista de fundos locais e os perfis das gestoras.

Uso:  python pipeline/extrair_guia.py "Guia-de-Fundos-Setembro-2026.xlsx"

Gera/atualiza, ao lado deste script:
  xp_fundos.csv   linhas "fundos" refeitas pela aba "Fundos XP Locais"; as linhas "previdencia" existentes são mantidas
                  (a planilha de previdência é outra e chega separada)
  perfis.json     abas ocultas "Base Assets" (gestoras), "Base Gestores" (carreira dos gestores), "Base Fundos" (estratégia,
                  equipe e gestores de cada fundo) e "Base Comentários" (atribuição de performance e posicionamento)

O pipeline (build_data.py) publica perfis.json em data/perfis/gestoras.json e data/perfis/fundos/<cnpj>.json.
"""
from __future__ import annotations

import csv
import datetime as dt
import json
import os
import re
import sys
import unicodedata

AQUI = os.path.dirname(os.path.abspath(__file__))


def norm(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def cnpj14(v) -> str:
    d = re.sub(r"\D", "", str(v or ""))
    return d.zfill(14) if d else ""


def texto(v) -> str:
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s in ("-", "None", "nan") else s


def ano(v):
    """Anos digitados como número viram datas de 1905 no Excel (serial = ano). Devolve o ano como texto, 'atual' ou ''."""
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        if v.year == 1905:
            return str((v - dt.datetime(1899, 12, 30)).days)
        return str(v.year)
    s = str(v).strip().lower()
    if s in ("", "-", "none", "nan"):
        return ""
    if s.startswith("atual"):
        return "atual"
    m = re.search(r"(19|20)\d{2}", s)
    return m.group(0) if m else s


def numero(v):
    try:
        if v is None or str(v).strip() in ("", "-"):
            return ""
        return float(v)
    except Exception:  # noqa: BLE001
        return ""


def linhas(ws, cabecalho: int):
    it = ws.iter_rows(min_row=cabecalho, values_only=True)
    hdr = [texto(h) for h in next(it)]
    for r in it:
        if all(c is None for c in r):
            continue
        yield dict(zip(hdr, r))


def extrair_locais(wb) -> list[dict]:
    """Aba 'Fundos XP Locais': o cabeçalho está na 4ª linha (índice de colunas na 1ª)."""
    ws = wb["Fundos XP Locais"]
    out = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        cnpj = cnpj14(r[1])
        if not cnpj or not r[2]:
            continue
        estrelas = str(r[5] or "").count("«")
        out.append({
            "cnpj": cnpj, "nome": texto(r[2]), "origem": "fundos", "cvm": texto(r[12]), "anbima": texto(r[13]),
            "classe_xp": texto(r[14]), "risco": texto(r[3]), "top": "1" if texto(r[4]).lower().startswith("s") else "0",
            "estrelas": str(estrelas) if estrelas else "", "gestor": texto(r[6]), "publico": texto(r[11]),
            "tributacao": texto(r[15]), "aplic_min": numero(r[16]), "status": texto(r[17]), "liquidez_dias": numero(r[25]),
            "taxa_adm": numero(r[29]), "taxa_perf": numero(r[32]), "benchmark": texto(r[35]), "pagina_xp": texto(r[37]),
            "ret12_xp": numero(r[40]), "pl_mil": numero(r[93]),
        })
    return out


CAMPOS_CSV = ["cnpj", "nome", "origem", "cvm", "anbima", "classe_xp", "risco", "top", "estrelas", "gestor", "publico",
              "tributacao", "aplic_min", "status", "liquidez_dias", "taxa_adm", "taxa_perf", "benchmark", "pagina_xp",
              "ret12_xp", "pl_mil"]


def atualizar_csv(locais: list[dict], caminho: str) -> tuple[int, int]:
    prev = []
    if os.path.exists(caminho):
        with open(caminho, encoding="utf-8") as fh:
            prev = [r for r in csv.DictReader(fh) if r.get("origem") == "previdencia"]
    with open(caminho, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=CAMPOS_CSV)
        w.writeheader()
        for r in locais:
            w.writerow({k: ("" if r.get(k) is None else r.get(k)) for k in CAMPOS_CSV})
        for r in prev:
            w.writerow({k: r.get(k, "") for k in CAMPOS_CSV})
    return len(locais), len(prev)


def extrair_perfis(wb) -> dict:
    # gestores: id -> nome + carreira
    gestores = {}
    for r in linhas(wb["Base Gestores"], 1):
        gid = texto(r.get("manager_id"))
        if not gid:
            continue
        carreira = []
        for k in range(1, 7):
            emp = texto(r.get(f"company_name_{k}"))
            if emp:
                carreira.append({"empresa": emp, "de": ano(r.get(f"company_since_{k}")), "ate": ano(r.get(f"company_until_{k}"))})
        gestores[gid] = {"nome": texto(r.get("manager_name")), "carreira": carreira}
    por_nome = {norm(g["nome"]): gid for gid, g in gestores.items()}

    # gestoras (assets): descrição + equipe principal
    gestoras = {}
    for r in linhas(wb["Base Assets"], 1):
        aid = texto(r.get("asset_id"))
        if not aid:
            continue
        equipe = []
        for k in range(1, 6):
            nome = texto(r.get(f"collab_name_{k}"))
            if nome:
                equipe.append({"nome": nome, "cargo": texto(r.get(f"collab_office_{k}")), "desde": ano(r.get(f"collab_market_entry_year_{k}")),
                               "gestor": por_nome.get(norm(nome))})
        gestoras[aid] = {"nome": texto(r.get("asset_name")), "descricao": texto(r.get("asset_description")), "equipe": equipe, "gestores": []}

    # comentários por CNPJ
    comentarios = {}
    for r in linhas(wb["Base Comentários"], 1):
        c = cnpj14(r.get("cnpj"))
        if not c:
            continue
        comentarios[c] = {"atribuicao": texto(r.get("atribuicao_de_performance")), "posicionamento": texto(r.get("posicionamento_atual")),
                          "pagina": texto(r.get("PAGINA_FUNDO_PLATAFORMA_XP")), "material": texto(r.get("material_promocional"))}

    # fundos
    fundos = {}
    for r in linhas(wb["Base Fundos"], 1):
        c = cnpj14(r.get("fund_cnpj"))
        if not c:
            continue
        aid = texto(r.get("asset_id"))
        ids = [texto(r.get(f"manager_{k}")) for k in range(1, 6)]
        ids = [i for i in ids if i and i in gestores]
        cats = [k.replace("category_", "") for k in r if str(k).startswith("category_") and texto(r.get(k)) in ("1", "1.0")]
        inicio = r.get("fund_date")
        fundos[c] = {
            "nome": texto(r.get("fund_name")), "gestora": aid if aid in gestoras else None,
            "inicio": inicio.strftime("%Y-%m-%d") if isinstance(inicio, dt.datetime) else texto(inicio)[:10],
            "estrategia": texto(r.get("fund_strategy")), "equipe": texto(r.get("fund_team")), "gestores": ids,
            "categorias": cats, "video": texto(r.get("youtube_video")),
        }
        if aid in gestoras:
            for i in ids:
                if i not in gestoras[aid]["gestores"]:
                    gestoras[aid]["gestores"].append(i)
    for c, com in comentarios.items():
        fundos.setdefault(c, {"nome": "", "gestora": None, "inicio": "", "estrategia": "", "equipe": "", "gestores": [], "categorias": [], "video": ""})
        fundos[c]["comentario"] = {k: v for k, v in com.items() if v}
    # gestores citados na equipe principal da gestora também entram na lista
    for g in gestoras.values():
        for m in g["equipe"]:
            if m.get("gestor") and m["gestor"] not in g["gestores"]:
                g["gestores"].append(m["gestor"])
    return {"gestoras": gestoras, "gestores": gestores, "fundos": fundos}


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 1
    import openpyxl

    wb = openpyxl.load_workbook(argv[1], read_only=True, data_only=True)
    locais = extrair_locais(wb)
    n_loc, n_prev = atualizar_csv(locais, os.path.join(AQUI, "xp_fundos.csv"))
    perfis = extrair_perfis(wb)
    perfis["fonte"] = os.path.basename(argv[1])
    perfis["gerado"] = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d")
    with open(os.path.join(AQUI, "perfis.json"), "w", encoding="utf-8") as fh:
        json.dump(perfis, fh, ensure_ascii=False, separators=(",", ":"))
    print(f"xp_fundos.csv: {n_loc} fundos locais novos + {n_prev} de previdência mantidos")
    print(f"perfis.json: {len(perfis['gestoras'])} gestoras, {len(perfis['gestores'])} gestores, {len(perfis['fundos'])} fundos "
          f"({sum(1 for f in perfis['fundos'].values() if f.get('comentario'))} com comentário)")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
